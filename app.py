from flask import Flask, request, jsonify
from flask_cors import CORS
import ingest_mod
import datetime
import os

app = Flask(__name__)
CORS(app)

@app.route('/ingest', methods=['POST'])
def ingest():
    try:
        data = request.json
        text = data.get('text')
        force = data.get('force', False)
        if not text:
            return jsonify({"status": "error", "message": "No se proporcionó texto"}), 400
        result = ingest_mod.process_text(text, force=force)
        return jsonify(result)
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/recent', methods=['GET'])
def recent():
    try:
        from openpyxl import load_workbook
        wb = load_workbook(ingest_mod.EXCEL_PATH, data_only=True)
        if 'Movimientos' not in wb.sheetnames: return jsonify([])
        sheet = wb['Movimientos']
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        recent_data = []
        today = datetime.date.today()
        yesterday = today - datetime.timedelta(days=1)
        
        # 1. CLEAN: Filter out totally empty rows (ghost rows)
        sheet = wb['Movimientos']
        all_rows = list(sheet.iter_rows(min_row=2, values_only=True))
        rows = [r for r in all_rows if r and r[0] is not None]
        
        # Calculate running balances consistently
        running_balances = []
        curr = 0
        for i, r in enumerate(rows):
            if len(r) > 7 and r[7] is not None and str(r[7]).strip() != "":
                try: 
                    val = float(str(r[7]).replace(',', '.'))
                    if val != 0: curr = val
                except: pass
            else:
                try: amt = float(str(r[1] or 0).replace(',', '.'))
                except: amt = 0
                tp = str(r[4]).lower() if r[4] else ""
                if tp == 'ingreso': curr += amt
                elif tp == 'gasto': curr -= amt
            running_balances.append(curr)

        recent_data = []
        # ... rest of the logic
        for i in range(len(rows) - 1, max(-1, len(rows) - 51), -1):
            row = rows[i]
            if not row[0]: continue
            dt = row[0]
            if not isinstance(dt, datetime.date):
                try: dt = datetime.datetime.strptime(str(dt).split(' ')[0], '%Y-%m-%d').date()
                except: dt = today
            
            group = "Anteriores"
            if dt == today: group = "Hoy"
            elif dt == yesterday: group = "Ayer"
            
            recent_data.append({
                "id": i + 2,
                "fecha": f"{dt.year}-{dt.month:02d}-{dt.day:02d}",
                "monto": row[1],
                "categoria": row[2],
                "detalle": row[3],
                "tipo": row[4],
                "tienda": row[5] if len(row) > 5 else "",
                "subcategoria": row[6] if len(row) > 6 else "",
                "saldo": round(running_balances[i], 2),
                "group": group
            })
        return jsonify(recent_data)
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/search', methods=['GET'])
def search():
    try:
        query = request.args.get('q', '').lower()
        from openpyxl import load_workbook
        wb = load_workbook(ingest_mod.EXCEL_PATH, data_only=True)
        if 'Movimientos' not in wb.sheetnames: return jsonify([])
        sheet = wb['Movimientos']
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        # 1. CLEAN: Filter out totally empty rows (ghost rows)
        sheet = wb['Movimientos']
        all_rows = list(sheet.iter_rows(min_row=2, values_only=True))
        rows = [r for r in all_rows if r and r[0] is not None]

        # Calculate running balances consistently
        running_balances = []
        curr = 0
        for i, r in enumerate(rows):
            if len(r) > 7 and r[7] is not None and str(r[7]).strip() != "":
                try: 
                    val = float(str(r[7]).replace(',', '.'))
                    if val != 0: curr = val
                except: pass
            else:
                try: amt = float(str(r[1] or 0).replace(',', '.'))
                except: amt = 0
                tp = str(r[4]).lower() if r[4] else ""
                if tp == 'ingreso': curr += amt
                elif tp == 'gasto': curr -= amt
            running_balances.append(curr)
        
        # Latest bank anchor found in the filtered rows
        latest_bank_saldo = 0
        anchor_idx = -1
        latest_date = datetime.date(1900, 1, 1)
        for i, r in enumerate(rows):
            if len(r) > 7 and r[7] is not None and str(r[7]).strip() != "":
                try:
                    val = float(str(r[7]).replace(',', '.'))
                    dt = r[0]
                    if not isinstance(dt, datetime.date):
                        try: dt = datetime.datetime.strptime(str(dt).split(' ')[0], '%Y-%m-%d').date()
                        except: continue
                    else:
                        if hasattr(dt, 'date'): dt = dt.date()
                    if dt >= latest_date and val != 0:
                        latest_date = dt
                        latest_bank_saldo = val
                        anchor_idx = i
                except: continue

        results = []
        for i, row in enumerate(rows):
            if not row[0]: continue
            match = False
            if not query: match = True
            else:
                cat = str(row[2] or '').lower()
                det = str(row[3] or '').lower()
                tnd = str(row[5] or '').lower()
                if query in cat or query in det or query in tnd:
                    match = True
            
            if match:
                dt = row[0]
                if not isinstance(dt, datetime.date):
                    try: dt = datetime.datetime.strptime(str(dt).split(' ')[0], '%Y-%m-%d').date()
                    except: dt = datetime.date.today()
                
                results.append({
                    "id": i + 2,
                    "fecha": f"{dt.year}-{dt.month:02d}-{dt.day:02d}",
                    "monto": row[1],
                    "categoria": row[2],
                    "detalle": row[3],
                    "tipo": row[4],
                    "tienda": row[5] if len(row) > 5 else "",
                    "subcategoria": row[6] if len(row) > 6 else "",
                    "saldo": round(running_balances[i], 2)
                })
        return jsonify(list(reversed(results)))
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/summary', methods=['GET'])
def get_summary():
    try:
        from openpyxl import load_workbook
        wb = load_workbook(ingest_mod.EXCEL_PATH, data_only=True)
        if 'Movimientos' not in wb.sheetnames: return jsonify({"total_income": 0, "total_expense": 0, "savings": 0})
        # 1. CLEAN: Filter out totally empty rows (ghost rows)
        sheet = wb['Movimientos']
        all_rows = list(sheet.iter_rows(min_row=2, values_only=True))
        rows = [r for r in all_rows if r and r[0] is not None]

        latest_bank_saldo = 0
        anchor_idx = -1
        latest_date = datetime.date(1900, 1, 1)
        
        # Chronological Anchor: Find the row with the LATEST DATE that has a saldo
        # If multiple rows on the same date have a saldo, the LAST ONE in the list (highest index) is the true latest.
        for i, r in enumerate(rows):
            if len(r) > 7 and r[7] is not None and str(r[7]).strip() != "":
                try:
                    val = float(str(r[7]).replace(',', '.'))
                    dt = r[0]
                    if not isinstance(dt, (datetime.datetime, datetime.date)):
                        try: dt = datetime.datetime.strptime(str(dt).split(' ')[0], '%Y-%m-%d').date()
                        except: continue
                    else:
                        if hasattr(dt, 'date'): dt = dt.date()
                    
                    # DT is >= or i is greater? We want the LAST row of the latest date.
                    if dt > latest_date:
                        latest_date = dt
                        latest_bank_saldo = val
                        anchor_idx = i
                    elif dt == latest_date:
                        # Same day, update anchor to this later row
                        latest_bank_saldo = val
                        anchor_idx = i
                except: continue
        
        # Final adjustment loop (total historical stats)
        # current_balance starts from the latest known bank anchor
        current_balance = latest_bank_saldo
        total_income, total_expense = 0, 0
        
        # To calculate historical stats correctly, we scan all rows.
        # To calculate current 'savings' (projected), we only add/sub movements AFTER the anchor.
        for i, row in enumerate(rows):
            try: amount = float(str(row[1] or 0).replace(',', '.'))
            except: amount = 0.0
            
            tipo = str(row[4]).lower() if row[4] else ""
            if tipo == 'ingreso': total_income += amount
            elif tipo == 'gasto': total_expense += amount
            
            if i > anchor_idx:
                if tipo == 'ingreso': current_balance += amount
                elif tipo == 'gasto': current_balance -= amount

        return jsonify({
            "total_income": round(total_income, 2),
            "total_expense": round(total_expense, 2),
            "savings": round(current_balance, 2),
            "bank_balance": round(latest_bank_saldo, 2), # New dedicated field
            "debug": {
                "rows_scanned": len(rows),
                "anchor_idx": anchor_idx,
                "anchor_value": latest_bank_saldo,
                "anchor_date": str(latest_date)
            }
        })
    except Exception as e: return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/analytics', methods=['GET'])
def analytics():
    try:
        data = ingest_mod.get_analytics()
        return jsonify(data)
    except Exception as e: return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/taxonomy', methods=['GET'])
def taxonomy():
    return jsonify(ingest_mod.get_taxonomy())

@app.route('/config', methods=['GET', 'POST'])
def config():
    if request.method == 'GET':
        return jsonify(ingest_mod.load_config())
    else:
        try:
            new_config = request.json
            with open(ingest_mod.CONFIG_PATH, 'w', encoding='utf-8') as f:
                json.dump(new_config, f, indent=4)
            return jsonify({"status": "success", "message": "Configuración guardada"})
        except Exception as e:
            return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/insights', methods=['GET'])
def insights():
    try:
        data = ingest_mod.get_analytics()
        # Simple AI insight generation
        prompt = f"""
        Eres un asesor financiero de ELITE. Analiza estos datos de gasto mensual y da un consejo breve (max 2 frases) y profesional.
        Datos: {json.dumps(data.get('current_month', {}), indent=2)}
        Comparativa: {data.get('mom_growth', 0)}% vs mes pasado.
        """
        response = ingest_mod.client.models.generate_content(
            model=ingest_mod.MODEL_NAME,
            contents=prompt
        )
        return jsonify({"insight": response.text.strip()})
    except Exception as e:
        return jsonify({"insight": "No se pudo generar el consejo en este momento."})

@app.route('/update', methods=['POST'])
def update():
    try:
        data = request.json
        row_id = int(data.get('id'))
        success, message = ingest_mod.update_transaction(row_id, data)
        return jsonify({"status": "success" if success else "error", "message": message})
    except Exception as e: return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/delete', methods=['POST'])
def delete():
    try:
        data = request.json
        row_id = int(data.get('id'))
        success, message = ingest_mod.delete_transaction(row_id)
        return jsonify({"status": "success" if success else "error", "message": message})
    except Exception as e: return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/bank/process', methods=['POST'])
def process_bank():
    try:
        if 'file' not in request.files:
            return jsonify({"status": "error", "message": "No hay archivo"}), 400
        file = request.files['file']
        if file.filename == '':
            return jsonify({"status": "error", "message": "Archivo sin nombre"}), 400
        
        # Save temp file
        import tempfile
        temp = tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx')
        file.save(temp.name)
        temp.close()
        
        result = ingest_mod.process_bank_statement(temp.name)
        os.unlink(temp.name)
        return jsonify(result)
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/ingest_data', methods=['POST'])
def ingest_data():
    try:
        data = request.json
        transactions = data.get('transactions', [])
        success, info = ingest_mod.append_to_excel(transactions, force=True)
        return jsonify(info)
    except Exception as e: return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/debug/clear-activity', methods=['POST'])
def debug_clear_activity():
    try:
        from openpyxl import load_workbook
        wb = load_workbook(ingest_mod.EXCEL_PATH)
        sheet = wb['Movimientos']
        # Delete all rows except header (row 1)
        if sheet.max_row > 1:
            sheet.delete_rows(2, sheet.max_row)
        wb.save(ingest_mod.EXCEL_PATH)
        return jsonify({"status": "success", "message": "Actividad borrada correctamente"}), 200
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

if __name__ == '__main__':
    app.run(port=5000, debug=True)
