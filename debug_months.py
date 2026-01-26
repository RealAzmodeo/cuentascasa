from openpyxl import load_workbook
from datetime import datetime, date
import os

EXCEL_PATH = 'd:/Proyectos/Antigravity Offline/Cuentas-Casa/Presupuesto-Anual-2025.xlsx'

def check_months():
    wb = load_workbook(EXCEL_PATH, data_only=True)
    sheet = wb['Movimientos']
    months = {}
    tx_count = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[8] == 'Gabi Revolut':
            tx_count += 1
            d = row[0]
            if isinstance(d, (datetime, date)):
                m = d.strftime('%Y-%m')
            else:
                m = str(d)[:7]
            months[m] = months.get(m, 0) + 1
    
    print(f"Total Gabi Revolut txs: {tx_count}")
    print("Months distribution:")
    for m in sorted(months.keys(), reverse=True):
        print(f"  {m}: {months[m]}")

if __name__ == '__main__':
    check_months()
