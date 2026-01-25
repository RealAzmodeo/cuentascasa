from openpyxl import load_workbook
import os
import datetime

EXCEL_PATH = "d:/Proyectos/Antigravity Offline/Cuentas-Casa/Presupuesto-Anual-2025.xlsx"

def inspect_excel():
    if not os.path.exists(EXCEL_PATH):
        print(f"Error: {EXCEL_PATH} not found.")
        return

    print(f"Inspecting: {EXCEL_PATH}")
    wb = load_workbook(EXCEL_PATH, data_only=True)
    if 'Movimientos' not in wb.sheetnames:
        print("Sheet 'Movimientos' not found.")
        return
    
    sheet = wb['Movimientos']
    print(f"Max row: {sheet.max_row}")
    
    # Read headers
    headers = [cell.value for cell in sheet[1]]
    print(f"Headers: {headers}")
    
    # Check for empty rows at the end or in between
    non_empty_rows = []
    for i, row in enumerate(sheet.iter_rows(min_row=2, values_only=True)):
        if row[0] is not None:
            non_empty_rows.append((i + 2, row))
    
    print(f"Non-empty data rows: {len(non_empty_rows)}")
    
    # Sample last 5 non-empty rows
    print("\nLast 5 transactions:")
    for idx, r in non_empty_rows[-5:]:
        print(f"Row {idx}: {r}")

if __name__ == "__main__":
    inspect_excel()
