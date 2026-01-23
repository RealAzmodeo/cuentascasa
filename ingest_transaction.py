import os
import sys
import json
import datetime
from google.generativeai import GenerativeModel, configure
from openpyxl import load_workbook

import time

# Configuration
API_KEY = "AIzaSyCSA0I0B_RTKbl-gL11xiNEnQ2vqIc4V1U"
EXCEL_PATH = 'd:/Proyectos/Antigravity Offline/Cuentas-Casa/Presupuesto-Anual-2025.xlsx'
CATEGORIES = ['Juguetes', 'Supermercado', 'Bizum', 'Juegos de Mesa', 'Restaurants', 'Videojuegos', 'Tarjeta Revolut', 'Tarjeta Credito', 'Inversiones', 'Zapatillas', 'Taxi', 'Stich', 'Otros']
MODEL_NAME = 'gemini-2.5-flash'

configure(api_key=API_KEY)
model = GenerativeModel(MODEL_NAME)

def extract_transaction(text, retries=3):
    prompt = f"""
    Eres un asistente contable experto. Tu tarea es extraer información de una descripción de gasto o ingreso.
    Entrada: "{text}"
    Fecha de hoy: {datetime.date.today()}
    
    Debes devolver un JSON válido con el siguiente formato:
    {{
        "fecha": "YYYY-MM-DD",
        "monto": 0.0,
        "categoria": "Una de estas: {', '.join(CATEGORIES)}",
        "detalle": "Breve descripción",
        "tipo": "Gasto" o "Ingreso"
    }}
    
    Reglas:
    - Si no hay fecha, usa la de hoy.
    - Si no encuentras una categoría exacta, usa 'Otros'.
    - El monto debe ser un número positivo.
    """
    
    for i in range(retries):
        try:
            response = model.generate_content(prompt)
            data = json.loads(response.text.replace('```json', '').replace('```', '').strip())
            return data
        except Exception as e:
            if "429" in str(e):
                print(f"Quota exceeded. Retrying in {5 * (i+1)} seconds...")
                time.sleep(5 * (i+1))
                continue
            print(f"Error parsing Gemini response: {e}")
            if hasattr(response, 'text'):
                print(f"Raw response: {response.text}")
            return None
    return None


def append_to_excel(data):
    wb = load_workbook(EXCEL_PATH)
    if 'Movimientos' not in wb.sheetnames:
        print("Error: Sheet 'Movimientos' not found.")
        return
    
    sheet = wb['Movimientos']
    sheet.append([
        data['fecha'],
        data['monto'],
        data['categoria'],
        data['detalle'],
        data['tipo']
    ])
    
    wb.save(EXCEL_PATH)
    print(f"Added: {data['monto']} in {data['categoria']} - {data['detalle']}")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python ingest_transaction.py 'description'")
        sys.exit(1)
    
    text_input = " ".join(sys.argv[1:])
    transaction_data = extract_transaction(text_input)
    
    if transaction_data:
        append_to_excel(transaction_data)
