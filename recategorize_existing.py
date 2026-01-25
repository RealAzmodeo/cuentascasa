import ingest_mod
from openpyxl import load_workbook
import os

EXCEL_PATH = ingest_mod.EXCEL_PATH

def recategorize_existing():
    if not os.path.exists(EXCEL_PATH):
        print("Excel not found")
        return

    wb = load_workbook(EXCEL_PATH)
    sheet = wb['Movimientos']
    
    updated_count = 0
    total_rows = sheet.max_row
    
    print(f"Analyzing {total_rows - 1} transactions for re-categorization...")
    
    for row_idx in range(2, total_rows + 1):
        # Columns: A:Fecha, B:Monto, C:Cat, D:Detalle, E:Tipo, F:Tienda, G:Subcat
        category = sheet.cell(row=row_idx, column=3).value
        detail = sheet.cell(row=row_idx, column=4).value
        t_type = sheet.cell(row=row_idx, column=5).value
        
        # Only re-categorize if it's currently 'Otros' or 'Varios' or empty
        if category in ['Otros', 'Varios', None, '']:
            match = ingest_mod.apply_rules(str(detail), t_type)
            if match:
                sheet.cell(row=row_idx, column=3).value = match['categoria']
                sheet.cell(row=row_idx, column=7).value = match['subcategoria']
                updated_count += 1
                # print(f"Updated Row {row_idx}: {detail} -> {match['categoria']}")

    if updated_count > 0:
        wb.save(EXCEL_PATH)
        print(f"Success: Updated {updated_count} transactions with new rules.")
    else:
        print("No transactions were updated (no matches found for 'Otros').")

if __name__ == "__main__":
    recategorize_existing()
