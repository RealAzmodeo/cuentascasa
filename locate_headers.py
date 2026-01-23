from openpyxl import load_workbook

file_path = 'd:/Proyectos/Antigravity Offline/Cuentas-Casa/Presupuesto-Anual-2025.xlsx'
wb = load_workbook(file_path, data_only=True)
sheet = wb['Gastos']

print("--- Header Identification ---")
for r in range(1, 15):
    row_vals = [sheet.cell(row=r, column=c).value for c in range(1, 16)]
    print(f"Row {r}: {row_vals}")

wb.close()
