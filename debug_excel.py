import openpyxl
import os
import datetime

EXCEL_PATH = 'd:/Proyectos/Antigravity Offline/Cuentas-Casa/Presupuesto-Anual-2025.xlsx'

def diagnose():
    if not os.path.exists(EXCEL_PATH):
        print(f"ERROR: File not found at {EXCEL_PATH}")
        return

    try:
        wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
        sheet = wb.active
        print(f"--- DIAGNOSIS for {EXCEL_PATH} ---")
        print(f"Sheet Name: {sheet.title}")
        print(f"Max Row: {sheet.max_row}")
        print(f"Max Col: {sheet.max_column}")
        
        print("\n--- FIRST 10 ROWS ---")
        for i, row in enumerate(sheet.iter_rows(max_row=10, values_only=True), 1):
            print(f"R{i}: {row}")
            
        print("\n--- NON-EMPTY ROW SCAN ---")
        rows_with_data = []
        for i, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if any(cell is not None for cell in row):
                rows_with_data.append(i)
                if len(rows_with_data) <= 5:
                    print(f"Found data at R{i}: {row}")
        
        print(f"Total rows with data (excluding header): {len(rows_with_data)}")
        if rows_with_data:
            print(f"First data row: {min(rows_with_data)}")
            print(f"Last data row: {max(rows_with_data)}")
            
            # Show the last 5 rows with data
            print("\n--- LAST 5 DATA ROWS ---")
            for r_idx in rows_with_data[-5:]:
                vals = [sheet.cell(row=r_idx, column=c).value for c in range(1, 9)]
                print(f"R{r_idx}: {vals}")

    except Exception as e:
        print(f"CRITICAL ERROR: {e}")

if __name__ == "__main__":
    diagnose()
