import os
import json
from google import genai
from dotenv import load_dotenv

load_dotenv()
API_KEY = "AIzaSyCSA0I0B_RTKbl-gL11xiNEnQ2vqIc4V1U"
client = genai.Client(api_key=API_KEY)

folder = r'd:\Proyectos\Antigravity Offline\Cuentas-Casa\Extractos bancarios'
files = os.listdir(folder)

results = {}

for filename in files:
    filepath = os.path.join(folder, filename)
    ext = os.path.splitext(filename)[1].lower()
    
    print(f"DEBUG: Analizando {filename}...")
    
    if ext == '.pdf':
        try:
            with open(filepath, 'rb') as f:
                uploaded_file = client.files.upload(file=f, config={'mime_type': 'application/pdf'})
            
            prompt = "Identifica el titular de la cuenta o el nombre de la cuenta en este extracto bancario. Devuelve solo el nombre."
            response = client.models.generate_content(model='gemini-2.5-flash', contents=[uploaded_file, prompt])
            results[filename] = response.text.strip()
        except Exception as e:
            results[filename] = f"Error PDF: {str(e)}"
            
    elif ext in ['.xlsx', '.xls']:
        try:
            from openpyxl import load_workbook
            wb = load_workbook(filepath, data_only=True)
            sheet = wb.active
            sample = []
            for row in sheet.iter_rows(max_row=10, values_only=True):
                sample.append([str(c) for c in row if c is not None])
            
            prompt = f"Analiza estas primeras filas de un Excel bancario e identifica el titular o el nombre de la cuenta. Devuelve solo el nombre.\n\nDatos: {json.dumps(sample)}"
            response = client.models.generate_content(model='gemini-2.5-flash', contents=prompt)
            results[filename] = response.text.strip()
        except Exception as e:
            results[filename] = f"Error Excel: {str(e)}"

print("\n--- RESULTADOS DE IDENTIFICACIÓN ---")
print(json.dumps(results, indent=2))
