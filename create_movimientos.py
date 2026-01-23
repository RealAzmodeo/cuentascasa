from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

file_path = 'd:/Proyectos/Antigravity Offline/Cuentas-Casa/Presupuesto-Anual-2025.xlsx'
wb = load_workbook(file_path)

if 'Movimientos' not in wb.sheetnames:
    sheet = wb.create_sheet('Movimientos', 0) # Create as the first sheet
    
    # Headers
    headers = ['Fecha', 'Monto', 'Categoría', 'Detalle', 'Tipo (Ingreso/Gasto)']
    sheet.append(headers)
    
    # Styling headers
    header_fill = PatternFill(start_color='4F81BD', end_color='4F81BD', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)
    
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        
    # Column widths
    sheet.column_dimensions['A'].width = 12
    sheet.column_dimensions['B'].width = 12
    sheet.column_dimensions['C'].width = 20
    sheet.column_dimensions['D'].width = 35
    sheet.column_dimensions['E'].width = 20

    print("Sheet 'Movimientos' created successfully.")
else:
    print("Sheet 'Movimientos' already exists.")

wb.save(file_path)
print(f"File saved: {file_path}")
