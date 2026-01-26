import os
import json
from openpyxl import load_workbook
from dotenv import load_dotenv

def run_health_check():
    print("--- Cuentas-Casa Health Check ---")
    load_dotenv()
    
    # 1. Check Files
    excel_path = os.getenv("EXCEL_PATH")
    taxonomy_path = "taxonomy.json"
    env_path = ".env"
    
    files_to_check = {
        "Environment file (.env)": env_path,
        "Excel Database": excel_path,
        "Taxonomy configuration": taxonomy_path
    }
    
    all_files_ok = True
    for name, path in files_to_check.items():
        if path and os.path.exists(path):
            print(f"[OK] {name} found at {path}")
        else:
            print(f"[ERROR] {name} MISSING at {path}")
            all_files_ok = False
            
    if not all_files_ok:
        return

    # 2. Check Excel Content
    try:
        wb = load_workbook(excel_path)
        print(f"[INFO] Sheets found: {wb.sheetnames}")
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            # Count non-empty rows (excluding header in row 1)
            count = 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                if any(cell is not None for cell in row):
                    count += 1
            print(f"|-- Sheet: {sheet_name:20} | Rows: {count}")
            
        if "Movimientos" in wb.sheetnames:
            ws = wb["Movimientos"]
            count = 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                if any(cell is not None for cell in row):
                    count += 1
            if count == 0:
                print("[WARNING] The 'Movimientos' sheet is empty!")
            else:
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if any(cell is not None for cell in row):
                        print(f"[INFO] Sample transaction in Movimientos: {row[0]} | {row[1]} | {row[2]}")
                        break
        else:
            print("[ERROR] Sheet 'Movimientos' NOT FOUND in Excel.")
    except Exception as e:
        print(f"[ERROR] Could not read Excel: {str(e)}")

    # 3. Check Taxonomy
    try:
        with open(taxonomy_path, "r", encoding="utf-8") as f:
            tax = json.load(f)
            categories = list(tax.get("taxonomy", {}).keys())
            accounts = list(tax.get("account_details", {}).keys())
            print(f"[OK] taxonomy.json is valid. Categories found: {len(categories)}. Accounts found: {len(accounts)}")
            print(f"[INFO] Accounts: {', '.join(accounts)}")
    except Exception as e:
        print(f"[ERROR] Could not read taxonomy.json: {str(e)}")

    print("--- Health Check Finished ---")

if __name__ == "__main__":
    run_health_check()
