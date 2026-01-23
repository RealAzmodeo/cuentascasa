import pandas as pd
from openpyxl import load_workbook

file_path = 'd:/Proyectos/Antigravity Offline/Cuentas-Casa/Presupuesto-Anual-2025.xlsx'

print(f"Analyzing: {file_path}")

# Get all sheet names
wb = load_workbook(file_path, read_only=True)
sheets = wb.sheetnames
print(f"Sheets found: {sheets}")

# Analyze each sheet
for sheet in sheets:
    print(f"\n--- Sheet: {sheet} ---")
    try:
        # Read with pandas to get a quick summary
        df = pd.read_excel(file_path, sheet_name=sheet, nrows=20)
        print("Headers:")
        print(df.columns.tolist())
        print("\nFirst 5 rows (sample):")
        print(df.head())
    except Exception as e:
        print(f"Could not read sheet {sheet}: {e}")

wb.close()
