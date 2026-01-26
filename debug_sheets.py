from openpyxl import load_workbook
import os
from dotenv import load_dotenv

load_dotenv()
excel_path = os.getenv("EXCEL_PATH")
wb = load_workbook(excel_path)
print("TOTAL SHEETS:", len(wb.sheetnames))
for s in wb.sheetnames:
    ws = wb[s]
    print(f"SHEET_NAME:[{s}] MAX_ROW:[{ws.max_row}]")
print("FINISHED")
