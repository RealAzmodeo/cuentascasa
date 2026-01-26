from openpyxl import load_workbook
from datetime import datetime, date

EXCEL_PATH = 'd:/Proyectos/Antigravity Offline/Cuentas-Casa/Presupuesto-Anual-2025.xlsx'

def check_types():
    wb = load_workbook(EXCEL_PATH, data_only=True)
    sheet = wb['Movimientos']
    types = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[8] == 'Gabi Revolut':
            d = row[0]
            m = d.strftime('%Y-%m') if isinstance(d, (datetime, date)) else str(d)[:7]
            if m == '2026-01':
                t = str(row[4]).strip().lower()
                types[t] = types.get(t, 0) + 1
    print(types)

if __name__ == '__main__':
    check_types()
