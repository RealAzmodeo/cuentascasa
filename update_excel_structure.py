from openpyxl import load_workbook
import os

path = r'd:\Proyectos\Antigravity Offline\Cuentas-Casa\Presupuesto-Anual-2025.xlsx'
try:
    wb = load_workbook(path)
    sheet = wb['Movimientos']
    
    # Check if headers already exist to avoid double adding
    if sheet.cell(row=1, column=9).value != 'Cuenta':
        sheet.cell(row=1, column=9).value = 'Cuenta'
        sheet.cell(row=1, column=10).value = 'LinkID'
        
        # Fill existing rows with default account
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row=row, column=1).value is not None:
                sheet.cell(row=row, column=9).value = 'Germán'
        
        wb.save(path)
        print("Excel structure updated successfully.")
    else:
        print("Excel structure already updated.")
except Exception as e:
    print(f"Error updating Excel: {e}")
