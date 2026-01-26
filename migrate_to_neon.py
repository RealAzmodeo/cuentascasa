import os
import datetime
from openpyxl import load_workbook
from dotenv import load_dotenv
from database import init_db, SessionLocal
from models import Transaction

load_dotenv()

EXCEL_PATH = os.getenv("EXCEL_PATH")

def migrate():
    print("--- Starting Migration to Neon.tech ---")
    
    # 1. Initialize DB (Create tables)
    init_db()
    
    # 2. Open Excel
    if not os.path.exists(EXCEL_PATH):
        print(f"ERROR: Excel file not found at {EXCEL_PATH}")
        return

    wb = load_workbook(EXCEL_PATH, data_only=True)
    if 'Movimientos' not in wb.sheetnames:
        print("ERROR: Sheet 'Movimientos' not found in Excel. Nothing to migrate.")
        return

    sheet = wb['Movimientos']
    db = SessionLocal()
    
    try:
        added_count = 0
        # Iterate rows starting from the second one (header is Row 1)
        for row in sheet.iter_rows(min_row=2, values_only=True):
            # Format: 0:Fecha, 1:Monto, 2:Categoría, 3:Detalle, 4:Tipo, 5:Tienda, 6:Subcategoría, 7:Saldo Banco, 8:Cuenta, 9:LinkID
            if not row[0] or row[1] is None:
                continue # Skip empty rows
            
            # Parse Date
            fecha = row[0]
            if isinstance(fecha, str):
                try:
                    fecha = datetime.datetime.strptime(fecha.split(' ')[0], '%Y-%m-%d').date()
                except:
                    print(f"Skipping row with invalid date: {row[0]}")
                    continue
            elif isinstance(fecha, datetime.datetime):
                fecha = fecha.date()
                
            new_tx = Transaction(
                fecha=fecha,
                monto=float(row[1]),
                categoria=str(row[2] or "Otros"),
                detalle=str(row[3] or ""),
                tipo=str(row[4] or "Gasto"),
                tienda=str(row[5] or ""),
                subcategoria=str(row[6] or "Varios"),
                saldo_banco=float(row[7]) if row[7] is not None else None,
                cuenta=str(row[8] or "Germán"),
                link_id=str(row[9] or "")
            )
            db.add(new_tx)
            added_count += 1
        
        db.commit()
        print(f"SUCCESS: Migrated {added_count} transactions to Neon.tech.")
        
    except Exception as e:
        print(f"ERROR during migration: {e}")
        db.rollback()
    finally:
        db.close()

if __name__ == "__main__":
    migrate()
