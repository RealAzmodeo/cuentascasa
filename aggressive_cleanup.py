import openpyxl
import os
import datetime

EXCEL_PATH = 'd:/Proyectos/Antigravity Offline/Cuentas-Casa/Presupuesto-Anual-2025.xlsx'

def aggressive_cleanup():
    print("--- AGGRESSIVE CLEANUP START ---")
    if not os.path.exists(EXCEL_PATH):
        print("Error: File not found.")
        return

    try:
        wb = openpyxl.load_workbook(EXCEL_PATH)
        sheet = wb['Movimientos']
        
        # 1. Read valid data
        data = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            date_val = row[0]
            if date_val is None: continue
            
            # Check if it's an empty-ish string
            if isinstance(date_val, str) and not date_val.strip():
                continue
                
            data.append(list(row))
        
        print(f"Detected {len(data)} valid data rows out of {sheet.max_row-1} total rows.")
        
        # 2. Re-create sheet perfectly
        wb.remove(sheet)
        new_sheet = wb.create_sheet('Movimientos', 0)
        new_sheet.append(['Fecha', 'Monto', 'Categoría', 'Detalle', 'Tipo', 'Tienda', 'Subcategoría', 'Saldo Banco'])
        
        for r in data:
            new_sheet.append(r)
            
        wb.save(EXCEL_PATH)
        print("--- AGGRESSIVE CLEANUP FINISHED: Excel is now clean. ---")
        
    except Exception as e:
        print(f"Error during cleanup: {e}")

if __name__ == "__main__":
    aggressive_cleanup()
